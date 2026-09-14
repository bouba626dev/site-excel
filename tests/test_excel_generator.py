import os
import tempfile
import unittest
from pathlib import Path

from excel_generator import generate_excel_file


class ExcelGeneratorTests(unittest.TestCase):
    def test_generate_excel_file_creates_workbook_with_requested_sheets(self):
        specification = {"activity": "test", "excel": {"sheets": ["Test", "Feuille 2"]}}

        with tempfile.TemporaryDirectory() as tmp_dir:
            original_base_dir = os.environ.get("EXCEL_GENERATOR_OUTPUT_DIR")
            os.environ["EXCEL_GENERATOR_OUTPUT_DIR"] = tmp_dir
            try:
                output_path = generate_excel_file(specification)
                self.assertTrue(Path(output_path).exists())
                self.assertTrue(output_path.endswith(".xlsx"))
                from openpyxl import load_workbook

                workbook = load_workbook(output_path)
                self.assertEqual(workbook.sheetnames, ["Test", "Feuille 2"])
                self.assertEqual(workbook["Test"]["A1"].value, "Test")
            finally:
                if original_base_dir is None:
                    os.environ.pop("EXCEL_GENERATOR_OUTPUT_DIR", None)
                else:
                    os.environ["EXCEL_GENERATOR_OUTPUT_DIR"] = original_base_dir


if __name__ == "__main__":
    unittest.main()
